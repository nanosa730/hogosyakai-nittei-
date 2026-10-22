"""デモ用の架空データ（Microsoft Forms の書き出し形式）と配布物テンプレートを作る。

実在の人物とは関係ありません。氏名は「生徒+学年組番号」の仮名です。
使い方: python3 demo/make_demo.py
"""
import datetime
import os
import random

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
BANDS = [
    "11月10日(火)13:30～14:30", "11月10日(火)14:30～15:45",
    "11月11日(水)13:30～14:30", "11月11日(水)14:30～15:45",
    "11月12日(木)13:30～14:30", "11月12日(木)14:30～15:45",
    "11月13日(金)13:30～14:30", "11月13日(金)14:30～15:45",
    "11月16日(月)13:30～14:30", "11月16日(月)14:30～15:45",
]
HEADER = ["ID", "開始時刻", "完了時刻", "メール", "名前", "最終変更時刻",
          "お子様の学級をお選びください。", "お子様の出席番号をお選びください。",
          "お子様の名前をご入力ください。", "懇談の希望日時について、どちらかをお選びください。",
          *BANDS, "兄弟姉妹がいる場合は、学年･組・名前を入力してください。"]
CLASSES = "ABCDE"
LETTER_NO = {c: i + 1 for i, c in enumerate(CLASSES)}


def name(g, c, n):
    return f"生徒{g}{LETTER_NO[c]}{n:02d}"


# 兄弟姉妹の書き方のゆれを含める（どちらの子の欄に書くか、書き方はばらばら）
SIBLINGS = {
    (3, "A", 3): "2年B組 生徒2203",
    (2, "B", 3): "3年A組　生徒3103",
    (3, "B", 5): "２－Ｃ 生徒２３０５",
    (3, "C", 8): "生徒2408（2年D組）",
    (2, "D", 8): "姉：3年C組・生徒3308",
    (3, "D", 10): "2年3組 生徒2310",          # 組を数字で書いている（3組＝C組）
    (3, "E", 12): "2年E組 生徒2512、1年A組 生徒1112",  # 1年は未読込 → 要確認
    (2, "A", 15): "3年 生徒3215",              # 組の書き忘れ（名前で見つかる）
    (3, "A", 20): "2年A組 せいと",             # 名前の誤り → 要確認
    (2, "C", 4): "弟 小学校4年 生徒9999",       # 他校 → 無視
    (3, "B", 30): "なし",
    (2, "E", 2): "2年E組 生徒2503",            # 同学年の別の子（双子・年子）
    (3, "C", 1): "2年A組 生徒2101",
    (3, "C", 2): "2年A組 生徒2102",
}


def make_grade(grade, seed):
    rnd = random.Random(seed)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(HEADER)
    rid = 1
    base = datetime.datetime(2026, 10, 1, 14, 0)
    for c in CLASSES:
        missing = set(rnd.sample(range(1, 36), 2))
        for n in range(1, 36):
            if n in missing:
                continue
            anytime = rnd.random() < 0.12
            marks = []
            for _ in BANDS:
                p = 0.0 if anytime else rnd.choice([0.2, 0.45, 0.6])
                marks.append("×" if rnd.random() < p else None)
            if not anytime and all(m is None for m in marks):
                marks[rnd.randrange(len(BANDS))] = "×"
            t0 = base + datetime.timedelta(minutes=rid * 97)
            row = [rid, t0, t0 + datetime.timedelta(minutes=2), "anonymous", None, None, c, f"{n:02d}",
                   name(grade, c, n), "いつでもよい" if anytime else "都合の悪い日時がある", *marks,
                   SIBLINGS.get((grade, c, n))]
            ws.append(row)
            rid += 1
    # 二重回答（後のほうが正）
    t1 = base + datetime.timedelta(days=8)
    ws.append([rid, t1, t1, "anonymous", None, None, "D", "01", name(grade, "D", 1),
               "都合の悪い日時がある", *(["×"] * 4 + [None] * 6), None])
    path = os.path.join(HERE, f"demo_懇談会希望調査_{grade}年.xlsx")
    wb.save(path)
    return path


def make_template():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "游明朝"
    st.font.size = Pt(11)
    p = doc.add_paragraph("令和８年１０月２９日")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph("保護者様")
    p = doc.add_paragraph("〇〇市立〇〇中学校")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph("個人懇談会の日時について（お知らせ）")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    doc.add_paragraph("　先日はご希望日時のご回答をいただき、ありがとうございました。"
                      "個人懇談会の日時が下記のとおり決まりましたので、お知らせいたします。")
    p = doc.add_paragraph("記")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 目印をわざと複数の書式に分けて入れる（Wordで編集したときと同じ状態を再現）
    p = doc.add_paragraph("１　生徒　")
    p.add_run("【学年】")
    p.add_run("年【組】組【番号】番　")
    r = p.add_run("【氏名】")
    r.bold = True
    p = doc.add_paragraph("２　日時　")
    r = p.add_run("【日")
    r.bold = True
    r = p.add_run("時】")
    r.bold = True
    doc.add_paragraph("３　場所　各学級の教室")
    doc.add_paragraph("　都合が悪くなった場合は、担任までご連絡ください。")
    path = os.path.join(HERE, "配布物テンプレート_見本.docx")
    doc.save(path)
    return path


if __name__ == "__main__":
    print(make_grade(2, 2))
    print(make_grade(3, 3))
    print(make_template())
